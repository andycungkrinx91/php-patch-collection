#!/usr/bin/env python3
"""
Regenerate the CVE-2026-109618 / CVE-2026-109619 coverage workbook.

Design contract (tokobetsu-jonin xlsx SOP):
  - strict light mode, no dark fills, no #000000
  - Calibri only
  - 4-colour accent ribbon from theme T02
  - header tint #EEF2F6, pearl zebra #F8FAFC
  - status pill tints: #DCFCE7 good, #E0E7FF tier, #FEF3C7 warn, #FEE2E2 crit
  - freeze panes + native AutoFilter on every data sheet
  - auto-fitted column widths

Data is the state actually observed on 2026-10-24. Scope is PHP 5.5 through 8.1;
branches above 8.1 are supported upstream and receive their fix from the release.
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = (
    "/home/andycungkrinx/experiment/portofolio/data/php-patch-collection/"
    "docs/test-report/cve-2026-109618-109619-coverage-matrix-2026-10-24.xlsx"
)

A1, A2, A3, A4 = "17B3A3", "3A8DDE", "8A4FD0", "E85DA0"
TITLE, SUB, HEAD_TINT, ZEBRA = "334155", "475569", "EEF2F6", "F8FAFC"
BORDER_C, BODY, MUTED = "CBD5E1", "334155", "64748B"
GOOD, TIER, WARN, CRIT = "DCFCE7", "E0E7FF", "FEF3C7", "FEE2E2"
GOOD_T, TIER_T, WARN_T, CRIT_T = "166534", "3730A3", "92400E", "991B1B"
PANEL = "F1F5F9"

THIN = Side(style="thin", color=BORDER_C)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def ribbon(ws, span=8):
    for i, c in enumerate([A1, A2, A3, A4]):
        ws.cell(row=1, column=1 + i, value=None).fill = PatternFill("solid", fgColor=c)
    for i in range(4, span):
        ws.cell(row=1, column=1 + i, value=None)


def title(ws, text, sub=None, span=8):
    ribbon(ws, span)
    ws.cell(row=2, column=1, value=text).font = Font(
        name="Calibri", size=16, bold=True, color=TITLE
    )
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    ws.cell(row=2, column=1).alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 26
    if sub:
        ws.cell(row=3, column=1, value=sub).font = Font(
            name="Calibri", size=10, color=MUTED
        )
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=span)


def kpi(ws, row, col, label, value, accent=A2):
    c = ws.cell(row=row, column=col, value=label)
    c.font = Font(name="Calibri", size=8, bold=True, color=MUTED)
    c.fill = PatternFill("solid", fgColor=PANEL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = BOX
    v = ws.cell(row=row + 1, column=col, value=value)
    v.font = Font(name="Calibri", size=16, bold=True, color=accent)
    v.alignment = Alignment(horizontal="center", vertical="center")
    v.border = Border(
        left=THIN, right=THIN, bottom=THIN, top=Side(style="medium", color=accent)
    )


def table(ws, top, headers, rows, widths, tint_col=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=top, column=j, value=h)
        c.font = Font(name="Calibri", size=9, bold=True, color=TITLE)
        c.fill = PatternFill("solid", fgColor=HEAD_TINT)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = Border(
            left=THIN, right=THIN, top=THIN, bottom=Side(style="medium", color=BORDER_C)
        )
    ws.row_dimensions[top].height = 22

    for i, r in enumerate(rows):
        rr = top + 1 + i
        for j, val in enumerate(r, start=1):
            c = ws.cell(row=rr, column=j, value=val)
            c.font = Font(name="Calibri", size=11, color=BODY)
            c.alignment = Alignment(vertical="center", wrap_text=(j == len(r)))
            c.border = BOX
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=ZEBRA)
        if tint_col:
            c = ws.cell(row=rr, column=tint_col)
            txt = str(c.value or "")
            if txt.startswith("NOT VULNERABLE") or txt in ("PASS", "VERIFIED", "FIXED"):
                c.fill = PatternFill("solid", fgColor=GOOD)
                c.font = Font(name="Calibri", size=11, bold=True, color=GOOD_T)
            elif txt.startswith("VULNERABLE") or txt == "FAIL":
                c.fill = PatternFill("solid", fgColor=CRIT)
                c.font = Font(name="Calibri", size=11, bold=True, color=CRIT_T)
            elif "INCONCLUSIVE" in txt or "BENCH" in txt or "CONTROL" in txt:
                c.fill = PatternFill("solid", fgColor=WARN)
                c.font = Font(name="Calibri", size=11, bold=True, color=WARN_T)
            elif txt in ("PATCHED", "FIXED RELEASE", "VENDOR"):
                c.fill = PatternFill("solid", fgColor=TIER)
                c.font = Font(name="Calibri", size=11, bold=True, color=TIER_T)

    # auto-fit
    for j, w in enumerate(widths, start=1):
        widest = len(str(headers[j - 1]))
        for i, _row in enumerate(rows):
            v = ws.cell(row=top + 1 + i, column=j).value
            widest = max(widest, len(str(v)) if v is not None else 0)
        ws.column_dimensions[get_column_letter(j)].width = max(w, min(widest + 3, 62))

    last = top + len(rows)
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{last}"
    return last + 2
